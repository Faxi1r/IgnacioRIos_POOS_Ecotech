import json
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from .validaciones import validar_fecha, validar_id, validar_texto


class Informe:
    """Guarda una copia de los datos usados al generar el reporte."""

    def __init__(self, id_informe, tipo, fecha_generacion=None, contenido="", creado_por=None):
        self.id_informe = validar_id(id_informe, opcional=True)
        self.tipo = tipo
        self.fecha_generacion = fecha_generacion or date.today()
        self.contenido = contenido
        self.creado_por = validar_id(creado_por, "El id del autor", opcional=True)

    @property
    def tipo(self):
        return self._tipo

    @tipo.setter
    def tipo(self, valor):
        self._tipo = validar_texto(valor, "El tipo de informe", 100)

    @property
    def fecha_generacion(self):
        return self._fecha_generacion

    @fecha_generacion.setter
    def fecha_generacion(self, valor):
        self._fecha_generacion = validar_fecha(valor, "La fecha del informe")

    @property
    def contenido(self):
        return self._contenido

    @contenido.setter
    def contenido(self, valor):
        if not isinstance(valor, str):
            raise ValueError("El contenido del informe debe ser texto.")
        self._contenido = valor

    def generar(self, empleados, departamentos, proyectos, registros) -> None:
        from .empleado import Empleado
        from .departamento import Departamento
        from .proyecto import Proyecto
        from .registro_tiempo import RegistroDeTiempo

        grupos = [(empleados, Empleado, "id_empleado"), (departamentos, Departamento, "id_departamento"),
                  (proyectos, Proyecto, "id_proyecto"), (registros, RegistroDeTiempo, "id_registro")]
        for lista, clase, campo in grupos:
            if not isinstance(lista, list) or any(not isinstance(x, clase) for x in lista):
                raise ValueError(f"Se necesita una lista de {clase.__name__}.")
            ids = [getattr(x, campo) for x in lista]
            if None in ids or len(ids) != len(set(ids)):
                raise ValueError("Use datos guardados, sin repetir el mismo id en una lista.")
            for elemento in lista:
                elemento.validar_datos()
        personas = {e.id_empleado: e for e in empleados}
        trabajos = {p.id_proyecto: p for p in proyectos}
        ids_departamentos = {d.id_departamento for d in departamentos}
        if any(e.id_departamento is not None and e.id_departamento not in ids_departamentos for e in empleados):
            raise ValueError("Faltan departamentos en los datos del informe.")
        if any(e.id_empleado not in personas for d in departamentos for e in d.empleados):
            raise ValueError("Faltan empleados de los departamentos.")
        if any(e.id_empleado not in personas for p in proyectos for e in p.empleados_asignados):
            raise ValueError("Faltan empleados de los proyectos.")
        if any(r.id_empleado not in personas or r.id_proyecto not in trabajos for r in registros):
            raise ValueError("Falta el empleado o proyecto de un registro.")
        datos = {
            "empleados": [
                [e.id_empleado, e.identificacion, e.nombre, e.correo, e.direccion, e.telefono,
                 e.fecha_inicio_contrato.isoformat(), e.salario, e.id_departamento]
                for e in empleados],
            "departamentos": [
                [d.id_departamento, d.nombre, d.gerente.nombre,
                 ", ".join(str(e.id_empleado) for e in d.empleados)] for d in departamentos],
            "proyectos": [
                [p.id_proyecto, p.nombre, p.descripcion, p.fecha_inicio.isoformat(),
                 ", ".join(str(e.id_empleado) for e in p.empleados_asignados)] for p in proyectos],
            "registros": [
                [r.id_registro, r.fecha.isoformat(), personas[r.id_empleado].nombre,
                 trabajos[r.id_proyecto].nombre, r.horas, r.descripcion_tareas]
                for r in sorted(registros, key=lambda x: (x.fecha, x.id_registro))],
        }
        self.fecha_generacion = date.today()
        self.contenido = json.dumps(datos, ensure_ascii=False)

    def _datos(self):
        try:
            datos = json.loads(self.contenido)
            columnas = {"empleados": 9, "departamentos": 4, "proyectos": 5, "registros": 6}
            if set(datos) != set(columnas):
                raise ValueError
            for clave, cantidad in columnas.items():
                if not isinstance(datos[clave], list) or any(
                        not isinstance(fila, list) or len(fila) != cantidad for fila in datos[clave]):
                    raise ValueError
            return datos
        except (ValueError, TypeError):
            raise ValueError("Genere un contenido válido antes de guardar o exportar el informe.") from None

    def _secciones(self):
        return [
            ("empleados", "Empleados", ["ID", "Identificación", "Nombre", "Correo", "Dirección", "Teléfono",
                                       "Contrato", "Salario", "Depto."]),
            ("departamentos", "Departamentos", ["ID", "Nombre", "Gerente", "Empleados (ID)"]),
            ("proyectos", "Proyectos", ["ID", "Nombre", "Descripción", "Inicio", "Empleados (ID)"]),
            ("registros", "Horas registradas", ["ID", "Fecha", "Empleado", "Proyecto", "Horas", "Tareas"]),
        ]

    def _texto_legible(self):
        datos = self._datos()
        lineas = [f"INFORME: {self.tipo}", f"Fecha: {self.fecha_generacion}",
                  f"Autor (id): {self.creado_por}",
                  f"Total de horas: {sum(r[4] for r in datos['registros']):.2f}"]
        for clave, titulo, columnas in self._secciones():
            lineas += ["", f"{titulo} ({len(datos[clave])})", " | ".join(columnas)]
            lineas += [" | ".join("" if v is None else str(v) for v in fila) for fila in datos[clave]]
        return "\n".join(lineas)

    def exportar_pdf(self, ruta: str) -> None:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        datos = self._datos()
        destino = Path(ruta)
        if destino.suffix.lower() != ".pdf":
            raise ValueError("El archivo debe terminar en .pdf.")
        destino.parent.mkdir(parents=True, exist_ok=True)
        fuente = Path(__file__).resolve().parents[2] / "assets" / "DejaVuSans.ttf"
        if "EcoTech" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("EcoTech", str(fuente)))
        estilos = getSampleStyleSheet()
        estilos["Heading1"].fontName = "EcoTech"
        estilos["Heading2"].fontName = "EcoTech"
        estilos.add(ParagraphStyle(name="Celda", fontName="EcoTech", fontSize=7, leading=10, wordWrap="CJK"))
        estilos.add(ParagraphStyle(name="TextoEco", fontName="EcoTech", fontSize=10, leading=15))
        doc = SimpleDocTemplate(str(destino), pagesize=landscape(A4), rightMargin=28, leftMargin=28,
                                topMargin=30, bottomMargin=30, title=self.tipo, author="EcoTech")
        ancho = landscape(A4)[0] - 56
        cuerpo = [Paragraph(escape(self.tipo), estilos["Heading1"]),
                  Paragraph(f"Fecha: {self.fecha_generacion} | Autor: {self.creado_por}", estilos["TextoEco"]),
                  Paragraph(f"Total de horas: {sum(r[4] for r in datos['registros']):.2f}", estilos["TextoEco"]),
                  Spacer(1, 12)]
        proporciones = {
            "empleados": [0.04, 0.10, 0.13, 0.16, 0.17, 0.11, 0.09, 0.12, 0.08],
            "departamentos": [0.06, 0.28, 0.30, 0.36],
            "proyectos": [0.06, 0.22, 0.40, 0.12, 0.20],
            "registros": [0.05, 0.11, 0.19, 0.20, 0.07, 0.38],
        }
        for clave, titulo, columnas in self._secciones():
            cuerpo.append(Paragraph(escape(titulo), estilos["Heading2"]))
            filas = [columnas] + datos[clave]
            celdas = [[Paragraph(escape("" if valor is None else str(valor)), estilos["Celda"])
                       for valor in fila] for fila in filas]
            tabla = Table(celdas, colWidths=[ancho * n for n in proporciones[clave]], repeatRows=1,
                          hAlign="LEFT", splitByRow=1, splitInRow=0)
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe9e4")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aebdb8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            cuerpo.extend([tabla, Spacer(1, 12)])
        def pie(canvas, documento):
            canvas.setFont("EcoTech", 8)
            canvas.drawRightString(landscape(A4)[0] - 28, 15, f"EcoTech - página {documento.page}")
        doc.build(cuerpo, onFirstPage=pie, onLaterPages=pie)

    def exportar_excel(self, ruta: str) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        datos = self._datos()
        destino = Path(ruta)
        if destino.suffix.lower() != ".xlsx":
            raise ValueError("El archivo debe terminar en .xlsx.")
        destino.parent.mkdir(parents=True, exist_ok=True)
        libro = Workbook()
        resumen = libro.active
        resumen.title = "Resumen"
        resumen.append(["Informe", self.tipo])
        resumen.append(["Fecha", self.fecha_generacion.isoformat()])
        resumen.append(["Autor (ID)", self.creado_por])
        resumen.append(["Total de horas", round(sum(r[4] for r in datos["registros"]), 2)])
        for clave, titulo, columnas in self._secciones():
            resumen.append([titulo, len(datos[clave])])
            hoja = libro.create_sheet(titulo)
            hoja.append(columnas)
            for fila in datos[clave]:
                hoja.append(fila)
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
        for hoja in libro:
            for fila in hoja:
                for celda in fila:
                    if isinstance(celda.value, str):
                        # Un texto que empieza con = debe seguir siendo texto.
                        celda.data_type = "s"
                    celda.alignment = Alignment(vertical="top", wrap_text=True)
            for celda in hoja[1]:
                celda.font = Font(bold=True, color="17372D")
                celda.fill = PatternFill("solid", fgColor="DBE9E4")
            for columna in hoja.columns:
                ancho = min(45, max(12, max(len(str(c.value or "")) for c in columna) + 2))
                hoja.column_dimensions[get_column_letter(columna[0].column)].width = ancho
            for fila in hoja:
                lineas = max(max(1, len(str(c.value or "")) // int(hoja.column_dimensions[c.column_letter].width) + 1)
                             for c in fila)
                hoja.row_dimensions[fila[0].row].height = min(300, 16 * lineas + 6)
        libro.save(destino)

    def __str__(self):
        return f"Informe(id={self.id_informe}, tipo='{self.tipo}', fecha={self.fecha_generacion}, autor={self.creado_por})"
